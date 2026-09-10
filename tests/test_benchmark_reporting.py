import os
import tempfile

import numpy as np
import pandas as pd
import pytest

from robustkit import (
    residual_summary, deviation_report, benchmark_report_suite, export_benchmark_excel,
    fit_huber_benchmark,
)


def make_gender_gap_df(n=800, seed=0):
    rng = np.random.default_rng(seed)
    employee_id = [f"E{i:04d}" for i in range(n)]
    age = rng.uniform(25, 60, n)
    gender = rng.choice(["M", "F"], n)
    job_family = rng.choice(["Eng", "Sales", "HR"], n, p=[0.5, 0.3, 0.2])
    location = rng.choice(["Stockholm", "Gothenburg"], n)
    salary = 30000 + 400 * age + rng.normal(0, 800, n)
    salary = salary - np.where(gender == "F", 1500, 0)  # injected gender gap

    return pd.DataFrame({
        "employee_id": employee_id, "age": age, "gender": gender,
        "job_family": job_family, "location": location, "salary": salary,
    })


# ---------------------------------------------------------------------------
# residual_summary
# ---------------------------------------------------------------------------

def test_residual_summary_detects_injected_gap():
    df = make_gender_gap_df()
    rs = residual_summary(df, segment_col="gender", y_col="salary", x_col="age")

    assert set(rs.columns) == {"segment", "n", "median_residual", "mad_residual", "p10_residual", "p90_residual"}
    assert rs.loc[rs["segment"] == "F", "median_residual"].iloc[0] < 0
    assert rs.loc[rs["segment"] == "M", "median_residual"].iloc[0] > 0
    assert (rs["mad_residual"] > 0).all()
    assert (rs["p90_residual"] > rs["p10_residual"]).all()


def test_residual_summary_requires_x_col_or_benchmark_fit():
    df = make_gender_gap_df()
    with pytest.raises(ValueError):
        residual_summary(df, segment_col="gender", y_col="salary")


# ---------------------------------------------------------------------------
# deviation_report (GAP #2)
# ---------------------------------------------------------------------------

def test_deviation_report_sorted_ascending_by_default():
    df = make_gender_gap_df()
    ndr = deviation_report(df, y_col="salary", x_col="age", top_n=20, id_cols=["employee_id", "gender"])

    assert len(ndr) == 20
    assert (ndr["residual"].diff().dropna() >= 0).all()
    assert set(ndr.columns) >= {"employee_id", "gender", "actual", "expected", "residual"}


def test_deviation_report_direction_positive_sorted_descending():
    df = make_gender_gap_df()
    ndr = deviation_report(df, y_col="salary", x_col="age", top_n=20, direction="positive")
    assert (ndr["residual"].diff().dropna() <= 0).all()


def test_deviation_report_direction_two_sided_sorted_by_magnitude():
    df = make_gender_gap_df()
    ndr = deviation_report(df, y_col="salary", x_col="age", top_n=20, direction="two_sided")
    abs_residuals = ndr["residual"].abs()
    assert (abs_residuals.diff().dropna() <= 0).all()


def test_deviation_report_rejects_invalid_direction():
    df = make_gender_gap_df()
    with pytest.raises(ValueError):
        deviation_report(df, y_col="salary", x_col="age", direction="sideways")


def test_deviation_report_reflects_injected_gap():
    """With a real injected gender gap, the most negative deviations
    should be disproportionately female."""
    df = make_gender_gap_df()
    ndr = deviation_report(df, y_col="salary", x_col="age", top_n=50, id_cols=["gender"])
    female_share = (ndr["gender"] == "F").mean()
    assert female_share > 0.7  # should be heavily skewed toward F given the injected -1500 gap


def test_deviation_report_falls_back_to_row_id_without_id_cols():
    df = make_gender_gap_df()
    ndr = deviation_report(df, y_col="salary", x_col="age", top_n=5)
    assert "row_id" in ndr.columns
    assert len(ndr) == 5


# ---------------------------------------------------------------------------
# benchmark_report_suite
# ---------------------------------------------------------------------------

def test_benchmark_report_suite_runs_all_group_columns():
    df = make_gender_gap_df()
    suite = benchmark_report_suite(
        df, group_columns=["gender", "job_family", "location"], y_col="salary", x_col="age", n_boot=100,
    )
    assert set(suite.keys()) == {"gender", "job_family", "location"}
    for report in suite.values():
        assert "difference" in report.columns


def test_benchmark_report_suite_reuses_same_benchmark_across_columns():
    """The benchmark should be fit once and reused, not refit per
    grouping column -- verified by passing a precomputed benchmark_fit
    and confirming results match calling segment_position_report
    directly with the same fit."""
    df = make_gender_gap_df()
    fit = fit_huber_benchmark(df["age"].to_numpy(dtype=float), df["salary"].to_numpy(dtype=float))

    suite = benchmark_report_suite(
        df, group_columns=["gender"], y_col="salary", benchmark_fit=fit, x_col="age", n_boot=100, seed=5,
    )

    from robustkit import segment_position_report
    direct = segment_position_report(df, segment_col="gender", y_col="salary", benchmark_fit=fit, x_col="age", n_boot=100, seed=5)

    pd.testing.assert_frame_equal(suite["gender"], direct)


# ---------------------------------------------------------------------------
# export_benchmark_excel
# ---------------------------------------------------------------------------

def test_export_benchmark_excel_creates_one_sheet_per_report():
    pytest.importorskip("openpyxl")  # skip gracefully if openpyxl isn't installed in this environment

    df = make_gender_gap_df()
    suite = benchmark_report_suite(
        df, group_columns=["gender", "job_family"], y_col="salary", x_col="age", n_boot=50,
    )

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        export_benchmark_excel(suite, path)

        assert os.path.exists(path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        assert set(wb.sheetnames) == {"gender", "job_family"}


def test_export_benchmark_excel_deduplicates_long_sheet_names():
    pytest.importorskip("openpyxl")  # skip gracefully if openpyxl isn't installed in this environment

    long_name_a = "a" * 40
    long_name_b = "a" * 39 + "b"  # truncates to the same 31 chars as long_name_a
    reports = {
        long_name_a: pd.DataFrame({"x": [1, 2]}),
        long_name_b: pd.DataFrame({"x": [3, 4]}),
    }

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "dedup.xlsx")
        export_benchmark_excel(reports, path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        assert len(wb.sheetnames) == 2
        assert len(set(wb.sheetnames)) == 2  # no collision
        assert all(len(name) <= 31 for name in wb.sheetnames)


def test_export_benchmark_excel_raises_clear_error_without_openpyxl(monkeypatch):
    """
    Regardless of whether openpyxl happens to be installed in the
    environment running this test, simulate its absence and confirm
    export_benchmark_excel raises a clear, actionable ImportError --
    the behavior an offline/air-gapped environment (no openpyxl, no
    way to install it) actually depends on.
    """
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("simulated: no module named 'openpyxl'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    with pytest.raises(ImportError, match="requires openpyxl"):
        export_benchmark_excel({"a": pd.DataFrame({"x": [1, 2]})}, "irrelevant_path.xlsx")
